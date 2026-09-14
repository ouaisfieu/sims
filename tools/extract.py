#!/usr/bin/env python3
"""
extract.py — Bureau federal du Plan workbook -> raw JSON.

Source : DATA_BUDGET2027_13320.xlsx, annexe de donnees du rapport
"Options de politiques pour l'elaboration du budget 2027" (BFP, 01/06/2026).

Ce script ne fait AUCUN jugement editorial : il transcrit la feuille telle
quelle, en nettoyant uniquement les espaces insecables et les blancs.
L'enrichissement editorial vit dans tools/enrichissement.json.

Usage : python3 tools/extract.py
Sortie : tools/source/mesures-brut.json
Requiert : openpyxl
"""
import json, re, unicodedata, pathlib, sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl requis : pip install openpyxl")

ROOT = pathlib.Path(__file__).resolve().parent.parent
XLSX = ROOT / "tools" / "source" / "DATA_BUDGET2027_13320.xlsx"
OUT  = ROOT / "tools" / "source" / "mesures-brut.json"

def clean(v):
    if v is None:
        return ""
    s = str(v).replace("\xa0", " ")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()

def slugify(s, maxlen=70):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-zA-Z0-9]+", "-", s).strip("-").lower()
    s = re.sub(r"-{2,}", "-", s)
    if len(s) > maxlen:
        s = s[:maxlen].rsplit("-", 1)[0]
    return s or "mesure"

PARTI_RE = re.compile(r"Maatregel\s+([A-Za-zÉé&\-\s]+?)\s*\d", re.U)
URL_RE   = re.compile(r"https?://\S+")

def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    out, seen = [], {}
    for i, r in enumerate(rows):
        cat, sous, desc, moda, motiv, montant, src = (list(r) + [None] * 7)[:7]
        desc = clean(desc)
        if not desc:
            continue

        src = clean(src)
        m = PARTI_RE.search(src)
        parti = m.group(1).strip() if m else None
        u = URL_RE.search(src)

        base = slugify(desc)
        n = seen.get(base, 0) + 1
        seen[base] = n
        slug = base if n == 1 else f"{base}-{n}"

        out.append({
            "ref": f"BFP-{i + 2:03d}",          # ligne reelle dans le classeur
            "slug": slug,
            "titre": desc,
            "categorie": clean(cat),
            "sousCategorie": clean(sous),
            "modalites": clean(moda),
            "motivation": clean(motiv),
            "montant": montant if isinstance(montant, (int, float)) else None,
            "parti": parti,
            "sourceTexte": src or None,
            "sourceUrl": u.group(0).rstrip(".,;") if u else None,
        })

    OUT.write_text(json.dumps({
        "source": {
            "fichier": XLSX.name,
            "publication": "Bureau federal du Plan, Options de politiques pour l'elaboration du budget 2027",
            "date": "2026-06-01",
            "url": "https://www.plan.be/sites/default/files/documents/REP_BUDGET2027_13320_FR.pdf",
            "unite": "millions d'euros par an",
        },
        "mesures": out,
    }, ensure_ascii=False, indent=1), encoding="utf-8")

    chiffrees = sum(1 for m in out if m["montant"] is not None)
    print(f"{len(out)} mesures ecrites dans {OUT.relative_to(ROOT)} ({chiffrees} chiffrees)")

if __name__ == "__main__":
    main()
